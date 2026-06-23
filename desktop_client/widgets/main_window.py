import json
import markdown
import time
import uuid

from PyQt6.QtWidgets import (
    QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QTextEdit, QTextBrowser, QPushButton, QComboBox,
    QLabel, QProgressBar, QTabWidget, QListWidget,
    QSplitter, QMessageBox, QFileDialog, QFrame, QSizePolicy,
    QScrollArea, QRadioButton, QButtonGroup,
    QTableWidget, QTableWidgetItem, QHeaderView, QAbstractItemView,
    QListWidgetItem
)
import re
from PyQt6.QtCore import Qt, QSize, QTimer, pyqtSignal, QThread
from desktop_client.worker import TaskWorker
from desktop_client.config import DEFAULT_USER_ID
from desktop_client.api_client import EduAgentApiClient
from desktop_client.widgets.prompt_edit import PromptInputCard
from desktop_client.widgets.history_item import HistoryItemWidget
from desktop_client.utils.markdown_renderer import MarkdownRenderer
from desktop_client.widgets.tutorial_view import TutorialView


from PyQt6.QtGui import QFontMetrics, QKeyEvent


class ResourceRenderWorker(QThread):
    finished_html = pyqtSignal(str, str)

    def __init__(self, key, value, references_md, md_renderer):
        super().__init__()
        self.key = key
        self.value = value
        self.references_md = references_md
        self.md_renderer = md_renderer

    def run(self):
        value = self.value
        if self.key == "mindmap":
            value = self.md_renderer.process_mermaid(value)
        content = value + self.references_md
        html_body = self.md_renderer.render_markdown(content)
        html = self.md_renderer.render_html_content(html_body)
        self.finished_html.emit(self.key, html)

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()

        self.setWindowTitle("EduAgentX 个性化学习多智能体系统")
        self.resize(1420, 860)

        self.worker = None
        self.current_task_id = None
        self.current_conversation_id = uuid.uuid4().hex
        self.latest_task = {}
        self.history_data = []
        self.global_profile = {}
        self.api_client = EduAgentApiClient()
        self.md_renderer = MarkdownRenderer()

        self.init_ui()
        self.check_backend_health(show_warning=True)
        
        self.health_timer = QTimer(self)
        self.health_timer.timeout.connect(lambda: self.check_backend_health(show_warning=False))
        self.health_timer.start(5000)

        # 实时耗时刷新定时器
        self.task_timer = QTimer(self)
        self.task_timer.timeout.connect(self.update_elapsed_time_tick)
        self.current_task_start_time = None
        self.loading_dots = 0

    def update_elapsed_time_tick(self):
        if self.current_task_start_time:
            elapsed = time.time() - self.current_task_start_time
            self.elapsed_label.setText(f"总耗时：{elapsed:.1f} 秒")
            
            # 让状态栏的省略号动起来，表明没有卡死
            self.loading_dots = (self.loading_dots + 1) % 4
            current_text = self.status_label.text()
            if "..." in current_text or "…" in current_text:
                base_text = current_text.replace("...", "").replace("…", "").strip()
                self.status_label.setText(f"{base_text}{'.' * self.loading_dots}")

    def showEvent(self, event) -> None:
        super().showEvent(event)

        if not getattr(self, "_initial_focus_set", False):
            self._initial_focus_set = True
            if hasattr(self, 'input_card'):
                QTimer.singleShot(0, self.input_card.input_box.setFocus)

    def set_task_running(self, running: bool, text: str = "") -> None:
        self.stop_button.setEnabled(running)
        self.mode_combo.setEnabled(not running)

        if hasattr(self, 'input_card'):
            self.input_card.set_processing_state(running, text)

        if not running and hasattr(self, 'input_card'):
            QTimer.singleShot(0, self.input_card.input_box.setFocus)

    def get_task_title(self, task: dict) -> str:
        title = str(task.get("task_title") or "").strip()

        if title:
            return title

        user_input = str(task.get("user_input") or "").strip()
        if user_input:
            return (
                user_input[:18] + "..."
                if len(user_input) > 18
                else user_input
            )

        task_id = str(task.get("task_id") or "")
        if task_id:
            return f"历史会话 {task_id[:8]}"

        return "历史会话"

    def reload_history_list(self):
        self.history_list.clear()

        if hasattr(self.history_list, "_widgets"):
            self.history_list._widgets.clear()

        for task in self.history_data:
            self.add_history_item_ui(task)

    def load_history(self):
        import os
        history_path = os.path.join(os.path.dirname(__file__), "..", "..", "data", "history.json")
        try:
            if os.path.exists(history_path):
                with open(history_path, "r", encoding="utf-8") as f:
                    self.history_data = json.load(f)
            else:
                self.history_data = []
        except Exception:
            self.history_data = []

        if self.history_data:
            for task in reversed(self.history_data):
                if task.get("student_profile"):
                    self.global_profile = dict(task.get("student_profile"))
                    self.update_profile(self.global_profile)
                    break
            
        self.reload_history_list()

    def add_history_item_ui(self, task: dict):
        title = self.get_task_title(task)
        task_id = str(task.get("task_id") or "")
        status = task.get("status", "completed")
        mode = task.get("mode", "概念讲解")
        is_current = (task_id == self.current_task_id)

        item = QListWidgetItem()
        item.setData(Qt.ItemDataRole.UserRole, task_id)
        item.setData(Qt.ItemDataRole.UserRole + 1, title)
        item.setSizeHint(QSize(0, 72))

        self.history_list.addItem(item)

        widget = HistoryItemWidget(title, status, mode, self.delete_history_item, item, is_current)
        self.history_list.setItemWidget(item, widget)

        if not hasattr(self.history_list, "_widgets"):
            self.history_list._widgets = []
        self.history_list._widgets.append(widget)

        return item
        
    def delete_history_item(self, item):
        row = self.history_list.row(item)
        if row < 0:
            return
            
        reply = QMessageBox.question(
            self, "确认删除", "确定要删除这条学习记录吗？删除后无法恢复。",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No
        )
        
        if reply != QMessageBox.StandardButton.Yes:
            return

        widget = self.history_list.itemWidget(item)
        self.history_list.takeItem(row)

        if hasattr(self.history_list, "_widgets") and widget:
            try:
                self.history_list._widgets.remove(widget)
            except ValueError:
                pass

        if widget:
            widget.deleteLater()

        if 0 <= row < len(self.history_data):
            deleted_task = self.history_data.pop(row)
            self.save_history()
                
            if self.latest_task.get("task_id") == deleted_task.get("task_id"):
                self.clear_result_area()
                self.latest_task = {}

    def save_history(self):
        import os
        data_dir = os.path.join(os.path.dirname(__file__), "..", "..", "data")
        os.makedirs(data_dir, exist_ok=True)
        history_path = os.path.join(data_dir, "history.json")
        try:
            with open(history_path, "w", encoding="utf-8") as f:
                json.dump(self.history_data, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print("保存历史记录失败:", e)

    def check_backend_health(self, show_warning=True):
        if not self.api_client.check_health():
            self.backend_status_label.setText("● 后端服务未连接")
            self.backend_status_label.setObjectName("BackendStatusDisconnected")
            self.backend_status_label.style().unpolish(self.backend_status_label)
            self.backend_status_label.style().polish(self.backend_status_label)
            if show_warning:
                QMessageBox.warning(
                    self, 
                    "后端服务未连接", 
                    "后端服务未启动，请先执行：\nuvicorn src.api.server:app --host 127.0.0.1 --port 8000"
                )
        else:
            self.backend_status_label.setText("● 后端服务已连接")
            self.backend_status_label.setObjectName("BackendStatusConnected")
            self.backend_status_label.style().unpolish(self.backend_status_label)
            self.backend_status_label.style().polish(self.backend_status_label)

    def make_card(self) -> QFrame:
        card = QFrame()
        card.setObjectName("Card")
        card.setFrameShape(QFrame.Shape.NoFrame)
        return card

    def make_section_title(self, text: str) -> QLabel:
        label = QLabel(text)
        label.setObjectName("SectionTitle")
        return label

    def make_subtitle(self, text: str) -> QLabel:
        label = QLabel(text)
        label.setObjectName("SectionSubTitle")
        label.setWordWrap(True)
        return label

    def init_ui(self):
        self.setWindowTitle("EduAgentX 个性化学习多智能体系统")
        self.resize(1420, 860)

        root = QWidget()
        root.setObjectName("RootWidget")

        root_layout = QVBoxLayout(root)
        root_layout.setContentsMargins(18, 18, 18, 18)
        root_layout.setSpacing(14)

        header = self.build_header_bar()
        root_layout.addWidget(header)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.setChildrenCollapsible(False)

        left_panel = self.build_left_panel()
        center_panel = self.build_center_panel()
        right_panel = self.build_right_panel()

        self.load_history()

        splitter.addWidget(left_panel)
        splitter.addWidget(center_panel)
        splitter.addWidget(right_panel)

        splitter.setStretchFactor(0, 2)
        splitter.setStretchFactor(1, 6)
        splitter.setStretchFactor(2, 2)

        root_layout.addWidget(splitter)

        self.setCentralWidget(root)

    def build_header_bar(self):
        header = QFrame()
        header.setObjectName("HeaderBar")
        header.setFixedHeight(74)

        layout = QHBoxLayout(header)
        layout.setContentsMargins(20, 10, 20, 10)
        layout.setSpacing(14)

        title_block = QWidget()
        title_layout = QVBoxLayout(title_block)
        title_layout.setContentsMargins(0, 0, 0, 0)
        title_layout.setSpacing(2)

        title = QLabel("EduAgentX")
        title.setObjectName("AppTitle")

        subtitle = QLabel("个性化学习多智能体系统")
        subtitle.setObjectName("AppSubtitle")

        title_layout.addWidget(title)
        title_layout.addWidget(subtitle)

        layout.addWidget(title_block)
        layout.addStretch()

        self.backend_status_label = QLabel("● 后端状态检测中")
        self.backend_status_label.setObjectName("BackendStatusDisconnected")
        layout.addWidget(self.backend_status_label)

        self.mode_combo = QComboBox()
        self.mode_combo.addItems(["快速模式", "完整模式", "答疑模式"])
        layout.addWidget(self.mode_combo)
        
        layout.addStretch()

        self.knowledge_manager_button = QPushButton("📚 资料库")
        self.knowledge_manager_button.setObjectName("KnowledgeManagerButton")
        self.knowledge_manager_button.clicked.connect(self.open_knowledge_manager)
        layout.addWidget(self.knowledge_manager_button)

        return header

    def open_knowledge_manager(self):
        from desktop_client.widgets.knowledge_manager import KnowledgeManagerDialog
        
        if getattr(self, "knowledge_manager", None) and self.knowledge_manager.isVisible():
            self.knowledge_manager.raise_()
            self.knowledge_manager.activateWindow()
            return

        self.knowledge_manager = KnowledgeManagerDialog(
            api_client=self.api_client,
            user_id=DEFAULT_USER_ID,
            conversation_id=self.current_conversation_id,
            parent=self
        )
        self.knowledge_manager.show()

    def build_left_panel(self):
        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(14)

        # 学生画像卡片
        profile_card = self.make_card()
        profile_layout = QVBoxLayout(profile_card)
        profile_layout.setContentsMargins(14, 14, 14, 14)
        profile_layout.setSpacing(10)

        profile_title = self.make_section_title("学生画像")
        profile_desc = self.make_subtitle("系统根据对话、练习与反馈动态更新")

        self.profile_browser = QTextBrowser()
        self.profile_browser.setObjectName("CardTextBrowser")
        self.profile_browser.setHtml(
            "<p style='color:#64748B;'>暂无学生画像。开始生成后将在这里展示。</p>"
        )

        profile_layout.addWidget(profile_title)
        profile_layout.addWidget(profile_desc)
        profile_layout.addWidget(self.profile_browser)

        layout.addWidget(profile_card, 3)

        # RAG 检索依据卡片
        rag_card = self.make_card()
        rag_layout = QVBoxLayout(rag_card)
        rag_layout.setContentsMargins(14, 14, 14, 14)
        rag_layout.setSpacing(10)
        rag_title = self.make_section_title("RAG 检索依据")
        self.source_list = QListWidget()
        self.source_list.setObjectName("SourceList")
        self.source_list.setStyleSheet("QListWidget { border: none; background-color: transparent; } QListWidget::item { padding: 4px; border-bottom: 1px solid #E2E8F0; }")
        self.source_list.setMaximumHeight(180)
        self.source_list.setMinimumHeight(120)
        rag_layout.addWidget(rag_title)
        rag_layout.addWidget(self.source_list)
        layout.addWidget(rag_card, 2)

        # 历史会话卡片
        history_card = self.make_card()
        history_layout = QVBoxLayout(history_card)
        history_layout.setContentsMargins(14, 14, 14, 14)
        history_layout.setSpacing(10)
        history_title = self.make_section_title("历史会话记录")
        self.history_list = QListWidget()
        self.history_list.setObjectName("HistoryList")
        self.history_list.setStyleSheet("QListWidget { border: none; background-color: transparent; } QListWidget::item { border-bottom: 1px solid #E2E8F0; } QListWidget::item:selected { background-color: #F1F5F9; color: #0F172A; border-radius: 4px; font-weight: bold; }")
        self.history_list.itemClicked.connect(self.on_history_clicked)
        history_layout.addWidget(history_title)
        history_layout.addWidget(self.history_list)
        layout.addWidget(history_card, 2)

        return panel

    def build_center_panel(self):
        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(14)

        # 输入卡片
        self.input_card = PromptInputCard()
        self.input_card.submitted.connect(self.start_task_from_card)
        self.input_card.demo_btn.clicked.connect(self.fill_demo_input)
        self.mode_combo.currentTextChanged.connect(self.input_card.update_mode_defaults)
        
        # New Session / Stop buttons go next to or around here... wait, PromptInputCard has start, demo, clear.
        # We need "新建会话", "取消任务", "导出结果" globally. 
        # Actually we can put them in a small top action bar above the input card or below it.
        # Let's put a global action bar below the input card.
        global_action_bar = QHBoxLayout()
        global_action_bar.setSpacing(8)
        
        self.new_session_button = QPushButton("新建学习会话")
        self.new_session_button.setObjectName("BlueButton")
        self.new_session_button.clicked.connect(self.start_new_session)
        
        self.stop_button = QPushButton("取消任务")
        self.stop_button.setObjectName("GrayButton")
        self.stop_button.clicked.connect(self.stop_polling)
        self.stop_button.setEnabled(False)
        
        self.export_button = QPushButton("导出结果")
        self.export_button.setObjectName("GreenButton")
        self.export_button.clicked.connect(self.export_result)
        
        global_action_bar.addWidget(self.new_session_button)
        global_action_bar.addWidget(self.stop_button)
        global_action_bar.addWidget(self.export_button)
        global_action_bar.addStretch()
        
        layout.addWidget(self.input_card)
        layout.addLayout(global_action_bar)

        # 资源 Tabs
        self.tabs = QTabWidget()
        self.tabs.currentChanged.connect(self.on_tab_changed)

        self.resource_views = {}
        self.tab_keys = []

        tab_names = {
            "doc": "讲义",
            "quiz": "练习",
            "plan": "计划",
            "answer": "答疑",
            "mindmap": "导图",
            "reading": "阅读",
            "code_case": "代码",
            "ppt": "PPT",
            "evaluation": "错题分析",
            "remedial": "查漏补缺"
        }

        empty_tips = {
            "doc": "课程讲义生成后将在这里展示。",
            "quiz": "练习题库生成后将在这里展示。",
            "plan": "学习计划生成后将在这里展示。",
            "answer": "智能答疑结果将在这里展示。",
            "mindmap": "思维导图生成后将在这里展示。",
            "reading": "拓展阅读生成后将在这里展示。",
            "code_case": "代码案例生成后将在这里展示。",
            "ppt": "PPT 大纲生成后将在这里展示。",
            "evaluation": "提交练习答案后，错题分析报告将在这里展示。",
            "remedial": "提交练习答案后，针对薄弱点的补充讲解和变式练习将在这里展示。"
        }

        for key, title in tab_names.items():
            if key in ["quiz", "plan"]:
                scroll = QScrollArea()
                scroll.setWidgetResizable(True)
                scroll.setObjectName(f"{key.capitalize()}ScrollArea")
                scroll.setStyleSheet(f"QScrollArea {{ border: none; background-color: transparent; }} QWidget#{key.capitalize()}Container {{ background-color: #FFFFFF; }}")
                
                container = QWidget()
                container.setObjectName(f"{key.capitalize()}Container")
                container_layout = QVBoxLayout(container)
                container_layout.setAlignment(Qt.AlignmentFlag.AlignTop)
                
                empty_label = QLabel(f"<div style='color:#64748B; padding:20px;'>{empty_tips[key]}</div>")
                container_layout.addWidget(empty_label)
                
                scroll.setWidget(container)
                self.resource_views[key] = scroll
                self.tabs.addTab(scroll, title)
            elif key == "doc":
                view = TutorialView()
                self.resource_views[key] = view
                self.tabs.addTab(view, title)
            else:
                browser = QTextBrowser()
                browser.setObjectName("ResourceBrowser")
                browser.setHtml(
                    f"<div style='color:#64748B; padding:20px;'>{empty_tips[key]}</div>"
                )
                self.resource_views[key] = browser
                self.tabs.addTab(browser, title)

            self.tab_keys.append(key)


        layout.addWidget(self.tabs, 1)

        return panel

    def build_right_panel(self):
        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(14)

        # 任务进度卡片
        progress_card = self.make_card()
        progress_layout = QVBoxLayout(progress_card)
        progress_layout.setContentsMargins(14, 14, 14, 14)
        progress_layout.setSpacing(10)

        progress_title = self.make_section_title("任务进度")

        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)

        self.status_label = QLabel("等待任务开始")
        self.status_label.setObjectName("SectionSubTitle")
        self.status_label.setWordWrap(True)

        self.elapsed_label = QLabel("总耗时：0.0 秒")
        self.elapsed_label.setObjectName("SectionSubTitle")

        progress_layout.addWidget(progress_title)
        progress_layout.addWidget(self.progress_bar)
        progress_layout.addWidget(self.status_label)
        progress_layout.addWidget(self.elapsed_label)

        # Agent 执行过程卡片
        agent_card = self.make_card()
        agent_layout = QVBoxLayout(agent_card)
        agent_layout.setContentsMargins(14, 14, 14, 14)
        agent_layout.setSpacing(10)

        agent_title = self.make_section_title("多智能体过程")
        self.agent_list = QListWidget()
        self.agent_list.addItem("等待智能体执行。")

        agent_layout.addWidget(agent_title)
        agent_layout.addWidget(self.agent_list)

        layout.addWidget(progress_card)
        layout.addWidget(agent_card, 1)

        return panel

    def start_new_session(self):
        self.current_conversation_id = uuid.uuid4().hex
        self.history_list.clearSelection()
        if hasattr(self, 'input_card'):
            self.input_card.clear_input()
            self.input_card.status_label.setVisible(False)
        self.clear_result_area()
        self.status_label.setText("新会话已就绪")
        self.global_profile = {}
        self.update_profile(self.global_profile)
        self.latest_task = {}

    def start_task_from_card(self, user_input: str):
        # 防止按钮点击、Enter 连按或信号重复触发。
        if self.worker is not None and self.worker.isRunning():
            return

        if not user_input:
            self.status_label.setText("请输入学习需求")
            self.input_card.input_box.setFocus()
            return

        mode = self.mode_combo.currentText()

        # 暂存本次提交内容，用于创建失败时恢复。
        self.last_submitted_text = user_input

        completed_history = [
            task for task in self.history_data 
            if task.get("status") in ["completed", "failed"]
            and task.get("conversation_id") == self.current_conversation_id
        ]
        history_for_api = [{"content": task.get("user_input", "")} for task in completed_history[-3:]]

        try:
            result = self.api_client.create_task(
                user_input=user_input,
                mode=mode,
                user_id="desktop_user",
                conversation_id=self.current_conversation_id,
                request_id=uuid.uuid4().hex,
                student_profile=dict(self.global_profile),
                history=history_for_api,
                requested_resources=self.input_card.get_requested_resources()
            )

            task_id = result.get("task_id")
            if not task_id:
                raise RuntimeError("后端未返回 task_id")

            self.current_task_id = task_id

            # 创建临时会话占位任务记录
            temp_task = {
                "task_id": task_id,
                "conversation_id": self.current_conversation_id,
                "user_input": user_input,
                "mode": mode,
                "status": "running",
                "progress": 0,
                "current_stage": "任务已提交，正在生成资源……",
                "agent_trace": [],
                "generated_resources": {},
                "retrieved_sources": [],
                "student_profile": dict(self.global_profile),
                "task_title": "新会话 (生成中...)"
            }
            self.history_data.append(temp_task)
            new_item = self.add_history_item_ui(temp_task)
            self.history_list.scrollToBottom()
            self.history_list.setCurrentItem(new_item)

            self.latest_task = temp_task

            # 仅在任务创建成功后清空和锁定。
            self.input_card.clear_input()
            self.set_task_running(True, user_input)

            self.status_label.setText("任务已提交，正在生成资源……")
            self.progress_bar.setValue(0)
            self.clear_result_area()

            self.current_task_start_time = time.time()
            self.task_timer.start(500)  # 每500毫秒更新一次动画和时间

            self.worker = TaskWorker(
                api_client=self.api_client,
                task_id=task_id,
            )
            self.worker.task_updated.connect(self.update_task_view)
            self.worker.task_finished.connect(self.on_task_finished)
            self.worker.task_failed.connect(self.on_task_failed)
            self.worker.start()

        except Exception as e:
            # 创建失败时保留用户原始输入。
            if not self.input_card.get_text().strip():
                self.status_label.setText(f"提交失败：{e}")
            self.set_task_running(False)
            self.input_card.set_text(self.last_submitted_text)
            QMessageBox.critical(self, "错误", f"提交任务失败：\n{e}")

    def stop_polling(self):
        if not self.current_task_id:
            return

        self.stop_button.setEnabled(False)
        self.status_label.setText("正在提交取消请求……")

        from desktop_client.worker import CancelWorker
        self.cancel_worker = CancelWorker(self.api_client, self.current_task_id)
        self.cancel_worker.cancel_success.connect(self.on_cancel_requested)
        self.cancel_worker.cancel_failed.connect(self.on_cancel_failed_ui)
        self.cancel_worker.start()

    def on_cancel_requested(self, response):
        self.status_label.setText(
            response.get(
                "message",
                "已提交取消请求，当前执行节点结束后将停止。",
            )
        )

    def on_cancel_failed_ui(self, error):
        self.stop_button.setEnabled(True)
        self.status_label.setText(f"取消请求提交失败：{error}")

    def update_task_view(self, task: dict, is_history_view: bool = False):
        task_id = task.get("task_id")

        # 1. 更新 self.history_data 内缓存的数据
        history_task = next((t for t in self.history_data if t.get("task_id") == task_id), None)
        if history_task:
            # 合并 student_profile 并写回 task 供后续使用
            profile_to_update = task.get("student_profile", {})
            if profile_to_update:
                for k, v in profile_to_update.items():
                    if v and v not in ["未知", "未识别", "未明确", "未说明"]:
                        self.global_profile[k] = v
                task["student_profile"] = dict(self.global_profile)
                
            history_task.update(task)

            # 同时更新左侧会话列表标题
            new_title = self.get_task_title(history_task)
            for i in range(self.history_list.count()):
                item = self.history_list.item(i)
                if item.data(Qt.ItemDataRole.UserRole) == task_id:
                    if item.data(Qt.ItemDataRole.UserRole + 1) != new_title:
                        item.setData(Qt.ItemDataRole.UserRole + 1, new_title)
                        widget = self.history_list.itemWidget(item)
                        if widget and hasattr(widget, "title_label"):
                            widget.title_label.setText(new_title)
                            widget.title_label.setToolTip(new_title)
                    break

        # 2. 判断用户当前是否正在查看此任务
        # 如果既不是查看历史视图，且当前查看的任务又不是此任务更新，则直接返回，不更新右侧主 UI
        if not is_history_view and self.latest_task.get("task_id") != task_id:
            return

        self.latest_task = task

        self.progress_bar.setValue(int(task.get("progress", 0)))
        
        new_stage = task.get("current_stage", "处理中...")
        resources = task.get("generated_resources", {})
        completed_count = sum(1 for v in resources.values() if v)
        total_count = self.tabs.count() if hasattr(self, 'tabs') else 8
        stats_text = f" ({completed_count} / {total_count} 个资源已完成)" if completed_count > 0 else ""

        if not self.status_label.text().startswith(new_stage.replace("...", "").replace("…", "")):
            self.status_label.setText(new_stage + stats_text)
        elif completed_count > 0:
            self.status_label.setText(new_stage + stats_text)

        created_at = task.get("created_at")
        updated_at = task.get("updated_at")
        if created_at and updated_at and is_history_view:
            total_time = updated_at - created_at
            self.elapsed_label.setText(f"总耗时：{total_time:.1f} 秒")

        if not is_history_view:
            self.update_profile(self.global_profile)
        else:
            self.update_profile(task.get("student_profile", {}))

        self.update_agent_trace(task.get("agent_trace", []))
        self.active_view_task_id = task.get("task_id", "")
        self.update_resources(task.get("generated_resources", {}), task.get("retrieved_sources", []))

    def update_profile(self, profile: dict):
        if not profile:
            self.profile_browser.setHtml(
                "<p style='color:#64748B;'>暂无学生画像。开始生成后将在这里展示。</p>"
            )
            return
            
        is_reused = hasattr(self, 'global_profile') and bool(self.global_profile) and self.global_profile == profile
        status_badge = '<span style="background-color:#E0F2FE; color:#0369A1; padding:2px 8px; border-radius:10px; font-size:11px; margin-left:8px;">已复用</span>' if is_reused else '<span style="background-color:#FEF3C7; color:#D97706; padding:2px 8px; border-radius:10px; font-size:11px; margin-left:8px;">已更新</span>'

        long_term_fields = {
            "major": "专业背景",
            "foundation": "知识基础",
            "cognitive_style": "认知风格",
            "learning_goal": "学习目标",
            "weakness": "知识短板"
        }
        
        short_term_fields = {
            "learning_pace": "学习节奏",
            "time_budget": "时间预算",
            "emotional_state": "情绪状态",
            "preferred_format": "资料偏好"
        }

        html = f"<div style='margin-bottom: 12px; font-weight: bold; color: #111827; border-bottom: 1px solid #E2E8F0; padding-bottom: 4px;'>长期学习画像 {status_badge}</div>"
        
        has_long = False
        for key, label in long_term_fields.items():
            value = profile.get(key)
            if value:
                has_long = True
                html += f"""
                <div style="margin-bottom:8px;">
                    <span style="color:#64748B; font-size: 12px;">{label}</span><br>
                    <b style="color:#0F172A; font-size: 13px;">{value}</b>
                </div>
                """
        if not has_long:
            html += "<p style='color:#64748B; font-size: 12px;'>暂无记录</p>"
            
        html += "<div style='margin-top: 16px; margin-bottom: 12px; font-weight: bold; color: #111827; border-bottom: 1px solid #E2E8F0; padding-bottom: 4px;'>本轮动态偏好</div>"
        
        has_short = False
        for key, label in short_term_fields.items():
            value = profile.get(key)
            if value:
                has_short = True
                html += f"""
                <div style="margin-bottom:8px;">
                    <span style="color:#64748B; font-size: 12px;">{label}</span><br>
                    <b style="color:#0F172A; font-size: 13px;">{value}</b>
                </div>
                """
        if not has_short:
            html += "<p style='color:#64748B; font-size: 12px;'>默认配置</p>"

        self.profile_browser.setHtml(html)

    def update_agent_trace(self, traces: list):
        self.agent_list.clear()

        if not traces:
            self.agent_list.addItem("等待智能体执行。")
            return

        for item in traces:
            status = item.get("status", "")
            agent = item.get("agent", "")
            message = item.get("message", "")
            latency = item.get("latency_ms")

            if status == "success":
                icon = "🟢"
            elif status == "failed":
                icon = "🔴"
            elif status == "warning":
                icon = "🟡"
            elif status == "running":
                icon = "⏳"
            else:
                icon = "⚪"

            latency_text = f" · {latency}ms" if latency else ""
            list_item = QListWidgetItem(f"{icon} {agent}\n{message}{latency_text}")
            self.agent_list.addItem(list_item)

    def update_resources(self, resources: dict, retrieved_sources: list = None):
        if retrieved_sources is None:
            retrieved_sources = []
            
        if hasattr(self, "source_list"):
            self.source_list.clear()
            if not retrieved_sources:
                self.source_list.addItem("暂无检索依据。")
            else:
                seen = set()
                for item in retrieved_sources:
                    source = item.get("source", "未知来源")
                    filename = source.replace('\\', '/').split('/')[-1]
                    if filename.endswith(".md"):
                        filename = filename[:-3]
                    elif filename.endswith(".pdf"):
                        filename = filename[:-4]
                        
                    page = item.get("page", "未知页码")
                    chunk_id = item.get("chunk_id", "")
                    preview = item.get("preview", "").replace("\n", " ")[:30] + "..."
                    
                    unique_key = f"{filename}_{page}_{chunk_id}"
                    if unique_key in seen:
                        continue
                    seen.add(unique_key)
                    
                    page_str = "" if page == "未知页码" else f" (第{page}页)"
                    list_item = QListWidgetItem(f"📄 {filename}{page_str}\n   [{chunk_id}] {preview}")
                    list_item.setToolTip(item.get("preview", ""))
                    self.source_list.addItem(list_item)
        
        references_md = ""
        if retrieved_sources:
            references_md = "\n\n---\n### 🔍 检索依据\n"
            seen = set()
            for item in retrieved_sources:
                source = item.get("source", "未知来源")
                filename = source.replace('\\', '/').split('/')[-1]
                if filename.endswith(".md"):
                    filename = filename[:-3]
                elif filename.endswith(".pdf"):
                    filename = filename[:-4]
                    
                page = item.get("page", "未知页码")
                chunk_id = item.get("chunk_id", "")
                preview = item.get("preview", "").replace("\n", " ")[:30] + "..."
                
                unique_key = f"{filename}_{page}_{chunk_id}"
                if unique_key in seen:
                    continue
                seen.add(unique_key)
                
                page_str = "" if page == "未知页码" else f" (第{page}页)"
                references_md += f"- **{filename}**{page_str} - `[{chunk_id}]` {preview}\n"

        self.raw_resources = resources
        self.references_md = references_md
        
        if not hasattr(self, "tab_keys") or not hasattr(self, "tabs"):
            return
        # Hide tabs without content
        for idx in range(self.tabs.count()):
            key = self.tab_keys[idx]
            has_content = bool(resources.get(key))
            self.tabs.setTabVisible(idx, has_content)
            
        current_idx = self.tabs.currentIndex()
        if current_idx >= 0:
            current_key = self.tab_keys[current_idx]
            self.render_resource(current_key)

    def render_resource(self, key: str):
        if not hasattr(self, "raw_resources") or not hasattr(self, "references_md"):
            return
            
        if not hasattr(self, '_resource_content_cache'):
            self._resource_content_cache = {}

        raw_value = self.raw_resources.get(key)
        if not raw_value or key not in self.resource_views:
            return
            
        if isinstance(raw_value, dict):
            status = raw_value.get("status", "completed")
            if status == "failed":
                value_str = f"生成失败: {raw_value.get('error', '未知错误')}"
            else:
                resource_data = raw_value.get("data", {})
                format_type = raw_value.get("format", "markdown")
                if format_type in ["markdown", "mermaid"]:
                    value_str = resource_data.get("content", "")
                elif format_type == "json":
                    value_str = json.dumps(resource_data, ensure_ascii=False)
                else:
                    value_str = str(resource_data)
        else:
            value_str = raw_value
            
        content_for_cache = value_str + self.references_md
        if self._resource_content_cache.get(key) == content_for_cache:
            return
            
        self._resource_content_cache[key] = content_for_cache

        if key == "doc":
            widgets_raw = self.raw_resources.get("doc_widgets", {})
            if isinstance(widgets_raw, dict) and "data" in widgets_raw:
                widgets = widgets_raw.get("data", {}).get("widgets", [])
            else:
                widgets = widgets_raw if isinstance(widgets_raw, list) else []
            task_id = getattr(self, "active_view_task_id", "")
            self.resource_views[key].load_resource(value_str, widgets, task_id=task_id)
            return

        if key == "quiz":
            self.md_renderer.render_quiz(value_str, self.references_md, self.resource_views[key].widget().layout())
            return
        if key == "plan":
            self.md_renderer.render_plan(value_str, self.references_md, self.resource_views[key].widget().layout())
            return
        
        if not hasattr(self, "_render_workers"):
            self._render_workers = []
            
        worker = ResourceRenderWorker(key, value_str, self.references_md, self.md_renderer)
        worker.finished_html.connect(self.on_render_finished)
        self._render_workers.append(worker)
        self._render_workers = [w for w in self._render_workers if not w.isFinished()]
        worker.start()
        
    def on_render_finished(self, key, html):
        if key in self.resource_views and hasattr(self.resource_views[key], 'setHtml'):
            self.resource_views[key].setHtml(html)
        
    def on_tab_changed(self, index):
        if hasattr(self, "tab_keys") and index >= 0 and index < len(self.tab_keys):
            self.render_resource(self.tab_keys[index])



    def closeEvent(self, event):
        super().closeEvent(event)

    def on_task_failed(self, error_message: str):
        self.task_timer.stop()
        if self.worker is not None:
            self.worker.stop()
            self.worker.wait(1500)
            self.worker = None

        task_id = self.current_task_id
        
        # 1. 从历史列表中彻底移除这个临时占位任务
        existing_index = next(
            (
                index
                for index, history_task
                in enumerate(self.history_data)
                if history_task.get("task_id") == task_id
            ),
            None,
        )
        if existing_index is not None:
            self.history_data.pop(existing_index)
            self.reload_history_list()
            self.save_history()

        # 2. 如果用户当前正在看着这个任务，重置主展示区并弹出提示
        if self.latest_task.get("task_id") == task_id:
            self.status_label.setText("任务执行失败")
            self.clear_result_area()
            self.latest_task = {}
            self.set_task_running(False)
            QMessageBox.warning(
                self,
                "任务执行失败",
                error_message or "任务执行过程中发生未知错误。",
            )
        else:
            # 否则，静默解锁运行状态
            self.set_task_running(False)

    def on_task_finished(self, task: dict):
        self.task_timer.stop()
        if self.worker is not None:
            self.worker.stop()
            self.worker.wait(1500)
            self.worker = None

        task_id = task.get("task_id")
        
        # 1. 查找并在 self.history_data 中更新/替换
        existing_index = next(
            (
                index
                for index, history_task
                in enumerate(self.history_data)
                if history_task.get("task_id") == task_id
            ),
            None,
        )

        if task.get("status") not in ["completed", "cancelled"]:
            # 任务执行失败：如果是在后台跑，我们直接从历史里移除它
            if existing_index is not None:
                self.history_data.pop(existing_index)
                self.reload_history_list()
                self.save_history()
            
            # 如果当前用户正看着它，显示失败提示且重置 UI
            if self.latest_task.get("task_id") == task_id:
                # 如果在后台跑，我们静默关闭锁定
                self.set_task_running(False)
                
                status = task.get("status")
                if status == "interrupted":
                    self.status_label.setText("服务异常中断，可重新执行")
                elif status == "failed":
                    self.status_label.setText("任务执行失败")
                elif status == "expired":
                    self.status_label.setText("任务已过期")
            return

        # 任务成功或取消：更新历史数据
        if existing_index is None:
            self.history_data.append(task)
            self.add_history_item_ui(task)
        else:
            self.history_data[existing_index] = task
            self.reload_history_list()

        self.save_history()
        
        # 2. 如果用户当前正看着此任务，更新主 UI
        if self.latest_task.get("task_id") == task_id:
            self.latest_task = task
            self.update_task_view(task)
            
            status = task.get("status")
            if status == "cancelled":
                self.status_label.setText("任务已取消")
            else:
                self.status_label.setText("生成完成")
                self.progress_bar.setValue(100)
                
            self.set_task_running(False)
            self.history_list.scrollToBottom()
        else:
            # 否则，静默重置全局锁定，使用户能够发起新生成
            self.set_task_running(False)

    def clear_result_area(self):
        self._resource_content_cache = {}
        self.agent_list.clear()
        if hasattr(self, "source_list"):
            self.source_list.clear()
        for key, view in self.resource_views.items():
            if key in ["quiz", "plan"]:
                container = view.widget()
                layout = container.layout()
                while layout.count():
                    child = layout.takeAt(0)
                    if child.widget():
                        child.widget().deleteLater()
                text = "练习题库生成后将在这里展示。" if key == "quiz" else "学习计划生成后将在这里展示。"
                empty_label = QLabel(f"<div style='color:#64748B; padding:20px;'>{text}</div>")
                layout.addWidget(empty_label)
            else:
                view.clear()

    def fill_demo_input(self):
        demo_text = (
            "我是计算机专业大二学生，正在学习人工智能导论。"
            "我的数学基础一般，尤其不理解梯度下降和反向传播。"
            "我希望通过图解、代码案例和练习题，在一周内掌握神经网络基础。"
        )
        self.input_card.set_text(demo_text)

    def on_history_clicked(self, item):
        task_id = item.data(Qt.ItemDataRole.UserRole)
        task = next((t for t in self.history_data if t.get("task_id") == task_id), None)
        if not task:
            QMessageBox.warning(self, "历史记录异常", "未找到对应的历史任务。")
            return

        self.current_conversation_id = task.get("conversation_id", uuid.uuid4().hex)

        is_active_running = (task_id == self.current_task_id and self.worker is not None and self.worker.isRunning())

        self.latest_task = task
        self.clear_result_area()
        self.update_task_view(task, is_history_view=True)

        self.set_task_running(is_active_running)

        title = self.get_task_title(task)
        if is_active_running:
            self.status_label.setText(task.get("current_stage", "正在生成……"))
        else:
            self.status_label.setText(f"正在查看：{title}")
        self.progress_bar.setValue(int(task.get("progress", 100)))

    def export_result(self):
        if not self.latest_task:
            QMessageBox.information(self, "提示", "暂无可导出的结果。")
            return

        path, filter_selected = QFileDialog.getSaveFileName(
            self,
            "导出结果",
            "eduagentx_result",
            "Markdown Files (*.md);;Word Documents (*.docx);;PDF Documents (*.pdf)"
        )

        if not path:
            return

        resources = self.latest_task.get("generated_resources", {})
        profile = self.latest_task.get("student_profile", {})
        sources = self.latest_task.get("retrieved_sources", [])
        review = self.latest_task.get("review_feedback", "")

        content = "# EduAgentX 学习资源生成结果\n\n"

        content += "## 学生画像\n\n"
        content += "```json\n"
        content += json.dumps(profile, ensure_ascii=False, indent=2)
        content += "\n```\n\n"

        content += "## RAG 检索参考依据\n\n"
        for item in sources:
            content += f"- {item.get('chunk_id', '')} | {item.get('source', '')} | 页码：{item.get('page', '')}\n"

        content += "\n## 生成资源\n\n"
        for key, value in resources.items():
            if key == "doc_widgets":
                content += "### 讲义互动练习\n\n"
                for i, w in enumerate(value, 1):
                    content += f"#### {i}. {w.get('title', '练习')}\n\n"
                    wtype = w.get("type")
                    if wtype == "inline_quiz":
                        content += f"**问题：** {w.get('question', '')}\n\n"
                        for opt in w.get("options", []):
                            content += f"- {opt.get('text', '')}\n"
                        content += f"\n**正确答案ID：** {w.get('correct_option_id', '')}\n\n"
                        content += f"**解析：** {w.get('explanation', '')}\n\n"
                    elif wtype == "sequence_sort":
                        content += "**待排序步骤：**\n\n"
                        for item in w.get("items", []):
                            content += f"- {item.get('text', '')}\n"
                        content += f"\n**参考顺序（ID）：** {', '.join(w.get('correct_order', []))}\n\n"
                    elif wtype == "matching_pairs":
                        content += "**匹配项：**\n\n"
                        for pair in w.get("pairs", []):
                            content += f"- {pair.get('left', '')} <=> {pair.get('right', '')}\n"
                        content += "\n"
                    elif wtype == "stepper_tutorial":
                        for step in w.get("steps", []):
                            content += f"**{step.get('title', '')}**: {step.get('content', '')} (提示: {step.get('hint', '')})\n\n"
            else:
                content += f"### {key}\n\n{value}\n\n"

        content += "## Reviewer 审核结果\n\n"
        content += review

        if filter_selected == "Markdown Files (*.md)" or path.endswith(".md"):
            with open(path, "w", encoding="utf-8") as f:
                f.write(content)
        elif filter_selected == "PDF Documents (*.pdf)" or path.endswith(".pdf"):
            self.export_to_pdf(content, path)
        elif filter_selected == "Word Documents (*.docx)" or path.endswith(".docx"):
            self.export_to_word(content, path)

        QMessageBox.information(self, "导出成功", f"结果已导出到：{path}")

    def export_to_pdf(self, markdown_content, path):
        from PyQt6.QtPrintSupport import QPrinter
        from PyQt6.QtGui import QTextDocument

        html = self.md_renderer.render_markdown(markdown_content)
        
        styled_html = f"""
        <html>
        <head>
        <style>
            body {{ font-family: 'Microsoft YaHei UI', sans-serif; line-height: 1.6; color: #333; }}
            h1, h2, h3 {{ color: #0F172A; }}
            table {{ border-collapse: collapse; width: 100%; margin-bottom: 1em; }}
            th, td {{ border: 1px solid #CBD5E1; padding: 8px; text-align: left; }}
            th {{ background-color: #F1F5F9; font-weight: bold; }}
            pre {{ background-color: #F8FAFC; padding: 12px; border: 1px solid #E2E8F0; border-radius: 4px; }}
            code {{ font-family: Consolas, monospace; }}
            blockquote {{ border-left: 4px solid #CBD5E1; margin: 0; padding-left: 12px; color: #64748B; }}
        </style>
        </head>
        <body>
        {html}
        </body>
        </html>
        """

        doc = QTextDocument()
        doc.setHtml(styled_html)
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(path)
        doc.print(printer)

    def export_to_word(self, markdown_content, path):
        try:
            import docx
            from docx.shared import Pt
        except ImportError:
            QMessageBox.warning(self, "警告", "尚未安装 python-docx 库，无法导出 Word。")
            return

        doc = docx.Document()
        
        lines = markdown_content.split('\n')
        in_code_block = False
        code_content = []
        
        for line in lines:
            if line.startswith('```'):
                if in_code_block:
                    in_code_block = False
                    p = doc.add_paragraph()
                    run = p.add_run('\n'.join(code_content))
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                    p.style = 'No Spacing'
                    code_content = []
                else:
                    in_code_block = True
            elif in_code_block:
                code_content.append(line)
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('- '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.strip() == '':
                continue
            else:
                doc.add_paragraph(line)
                
        doc.save(path)
